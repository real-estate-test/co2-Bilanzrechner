#!/usr/bin/env python3
# =====================================================================
#  Baut projektrechner-vorlage.xlsx
#
#  Die Vorlage geht an den Architekten. Sie liegt als fertige Datei im
#  Repository, damit sie sich ohne Werkzeug weitergeben lässt — dieses
#  Skript ist ihre Quelle. Eine Binärdatei ohne Quelle liesse sich
#  weder nachvollziehen noch sauber ändern.
#
#  WICHTIG: Die Spaltenüberschriften stehen zugleich in
#  app/import-excel.js (SP_FLAECHEN, SP_KUBATUR, SP_SPIEGEL). Wer hier
#  eine ändert, muss sie dort mitändern — sonst findet der Import die
#  Spalte nicht mehr. Die Selbsttests in tests/engine.html prüfen die
#  Namen gegeneinander.
#
#  Ausführen:  python3 vorlagen/vorlage-bauen.py
# =====================================================================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

KOPF  = PatternFill('solid', fgColor='232C39')
KOPFS = Font(color='FFFFFF', bold=True, size=10)
TITEL = Font(bold=True, size=14)
FETT  = Font(bold=True, size=11)
HINW  = Font(italic=True, color='6B7484', size=9)
EING  = PatternFill('solid', fgColor='FFFDF0')
RAND  = Border(*[Side('thin', color='D5DAE2')] * 4)


def kopfzeile(ws, zeile, spalten):
    for i, (titel, breite) in enumerate(spalten, start=1):
        c = ws.cell(row=zeile, column=i, value=titel)
        c.fill, c.font = KOPF, KOPFS
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = RAND
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.row_dimensions[zeile].height = 28


def eingabe(ws, von, bis, spalten):
    for r in range(von, bis + 1):
        for c in range(1, spalten + 1):
            z = ws.cell(row=r, column=c)
            z.fill, z.border = EING, RAND


wb = Workbook()

# ---------------------------------------------------------------- Anleitung
ws = wb.active
ws.title = 'Anleitung'
ws.column_dimensions['A'].width = 100
for i, (t, f) in enumerate([
    ('Projektrechner — Vorlage für Flächen, Kubaturen und Wohnungsspiegel', TITEL),
    ('', None),
    ('Diese Datei ausfüllen und dem Projektentwickler zurückgeben. Er liest sie', None),
    ('direkt ein — je vollständiger sie ist, desto weniger wird von Hand übertragen', None),
    ('und desto weniger Fehler entstehen dabei.', None),
    ('', None),
    ('So gehen Sie vor', FETT),
    ('1. Nur die hellgelben Felder ausfüllen. Die Spaltenköpfe bitte NICHT ändern,', None),
    ('   umbenennen oder verschieben — daran erkennt das Werkzeug die Spalten.', None),
    ('2. Zeilen dürfen Sie hinzufügen, besonders im Wohnungsspiegel.', None),
    ('3. Was Sie nicht wissen, lassen Sie leer. Leere Felder werden übersprungen;', None),
    ('   der vorhandene Wert im Werkzeug bleibt dann stehen.', None),
    ('4. Zahlen als Zahlen eintragen, nicht als Text. Einheiten stehen im Kopf', None),
    ('   und gehören nicht in die Zelle.', None),
    ('5. Die Datei am Ende aus Excel heraus als .xlsx speichern.', None),
    ('', None),
    ('Die drei Blätter', FETT),
    ('Flächen          je Objekt eine Zeile — Geschossfläche, Geschosse, Untergeschoss, Parkplätze', None),
    ('Kubaturen        je Objekt eine Zeile — Geschosshöhen oder Volumen', None),
    ('Wohnungsspiegel  je Wohnung eine Zeile — Nr., Geschoss, Zimmer, Fläche, Preis oder Miete', None),
    ('', None),
    ('Zu den Objekten', FETT),
    ('Neubau       was neu gebaut wird', None),
    ('Bestand      was stehen bleibt und saniert wird', None),
    ('Erweiterung  Anbau oder Aufstockung am Bestand', None),
    ('', None),
    ('Nicht zutreffende Objekte einfach leer lassen.', HINW),
], start=1):
    c = ws.cell(row=i, column=1, value=t)
    if f:
        c.font = f

# ---------------------------------------------------------------- Flächen
ws = wb.create_sheet('Flächen')
ws.cell(row=1, column=1, value='Flächen je Objekt').font = TITEL
ws.cell(row=2, column=1,
        value='Nur ausfüllen, was aus Ihrer Studie hervorgeht. Leere Felder bleiben '
              'im Werkzeug unverändert.').font = HINW
kopfzeile(ws, 4, [
    ('Objekt', 16), ('Geschossfläche oberirdisch m²', 18), ('Vollgeschosse', 13),
    ('Untergeschoss m²', 15), ('Nutzfläche NWF m²', 16), ('Parkplätze Stk.', 13),
    ('Fläche je Parkplatz m²', 16),
])
for i, name in enumerate(['Neubau', 'Bestand', 'Erweiterung'], start=5):
    ws.cell(row=i, column=1, value=name).font = Font(bold=True)
eingabe(ws, 5, 7, 7)
for i, t in enumerate([
    'Geschossfläche oberirdisch = Summe aller Vollgeschosse, ohne Untergeschoss.',
    'Untergeschoss ohne Einstellhalle — die wird über die Parkplätze gerechnet.',
    'Nutzfläche leer lassen, wenn sie sich aus der Quote ergeben soll.',
], start=9):
    ws.cell(row=i, column=1, value=t).font = HINW

# ---------------------------------------------------------------- Kubaturen
ws = wb.create_sheet('Kubaturen')
ws.cell(row=1, column=1, value='Kubaturen je Objekt').font = TITEL
ws.cell(row=2, column=1,
        value='Entweder die Höhen ODER die Volumen ausfüllen. Sind Volumen '
              'eingetragen, haben sie Vorrang.').font = HINW
kopfzeile(ws, 4, [
    ('Objekt', 16), ('Regelgeschosshöhe m', 15), ('Dachgeschosshöhe m', 15),
    ('Höhe Untergeschoss m', 16), ('Höhe Einstellhalle m', 16),
    ('Volumen oberirdisch m³', 17), ('Volumen Untergeschoss m³', 18),
    ('Volumen Einstellhalle m³', 18),
])
for i, name in enumerate(['Neubau', 'Bestand', 'Erweiterung'], start=5):
    ws.cell(row=i, column=1, value=name).font = Font(bold=True)
eingabe(ws, 5, 7, 8)
ws.cell(row=9, column=1, value='Volumen nach SIA 416, Aussenmass.').font = HINW

# ---------------------------------------------------------------- Wohnungsspiegel
#  Spaltennamen und Reihenfolge folgen der Tabelle im Werkzeug, damit
#  man beim Übertragen nicht umdenken muss. «Nr.» ist die Bezeichnung
#  der Wohnung und darf Text sein (A 1.1), «Geschoss» ist eine Zahl.
ws = wb.create_sheet('Wohnungsspiegel')
ws.cell(row=1, column=1, value='Wohnungsspiegel').font = TITEL
ws.cell(row=2, column=1,
        value='Je Wohnung eine Zeile. Zeilen dürfen Sie beliebig viele '
              'hinzufügen.').font = HINW
kopfzeile(ws, 4, [
    ('Haus', 14), ('Verwertung', 13), ('Nr.', 12), ('Geschoss', 10),
    ('Zimmer', 9), ('Fläche m²', 11), ('Preis CHF', 13),
    ('Miete CHF/Monat', 15), ('Anzahl', 9),
])
for i, zeile in enumerate([
    ('Haus A', 'STWE',  'A 1.1', 0, 3.5,  92,  850000, None, 1),
    ('Haus A', 'STWE',  'A 1.2', 0, 4.5, 118, 1050000, None, 1),
    ('Haus B', 'Miete', 'B 2.1', 1, 2.5,  64,    None, 1650, 1),
], start=5):
    for j, w in enumerate(zeile, start=1):
        if w is not None:
            ws.cell(row=i, column=j, value=w)
eingabe(ws, 5, 60, 9)
for i, t in enumerate([
    'Verwertung: STWE (Verkauf) oder Miete. Preis nur bei STWE, Miete nur bei Miete.',
    'Nr.: die Bezeichnung der Wohnung, z. B. A 1.1 — Text ist erlaubt.',
    'Geschoss: 0 = Erdgeschoss, 1 = 1. Obergeschoss, -1 = Untergeschoss.',
    'Anzahl: wie oft diese Wohnung vorkommt. Leer oder 1 = einmal.',
    'Die drei Zeilen oben sind Beispiele — bitte überschreiben oder löschen.',
], start=62):
    ws.cell(row=i, column=1, value=t).font = HINW

ziel = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                    'projektrechner-vorlage.xlsx')
wb.save(ziel)
print('gebaut:', ziel)
print('Blätter:', ', '.join(wb.sheetnames))
